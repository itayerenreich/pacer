"""From a file to chunks: extraction, cleaning, chunking strategies, the manifest, and import.

The sample file is built in the test (python-docx), so the tests need no file of yours.
"""

import json
from pathlib import Path

import docx
import pytest
from fake_llm import FakeEmbedder
from fakes import FakeUnitOfWork, empty_repos

from pacer.ai.embeddings import for_embedding
from pacer.ai.library import index_library, search_library
from pacer.domain import EvidenceLevel
from pacer.evaluation import load_library_set, unknown_sources
from pacer.library.chunking import SUMMARY, Unit, fixed, label_of, structure
from pacer.library.clean import clean_blocks, clean_text
from pacer.library.extract import Block, ExtractionError, extract
from pacer.library.ingest import ManifestError, import_file, plan, preview, split_sections
from pacer.library.manifest import Manifest

RAMBAM = [
    "1. סדרי החינוך: רמב\"ם, הלכות תלמוד תורה פרק ב'",
    "תקציר: פרק על חובות הציבור והמורה ומספר התלמידים בכיתה.",
    '"הלכה א: מושיבין מלמדי תינוקות בכל מדינה ומדינה ובכל פלך ופלך.',
    "הלכה ב: מכניסין את התינוקות להתלמד כבן שש כבן שבע לפי כח הבן ובנין גופו.",
    "הלכה ה: עשרים וחמשה תינוקות למלמד אחד, היו יתר על עשרים וחמשה עד ארבעים מושיבין עמו אחר לסייעו.\"",
    "2. חנוך לנער: משלי",
    "הפסוק המרכזי על חינוך אישי לפי טבע הילד.",
    'פסוק: "חֲנֹךְ לַנַּעַר עַל פִּי דַרְכּוֹ."',
]


def manifest(**changes) -> Manifest:
    data = {
        "file": "sample.docx",
        "split": "numbered",
        "sources": [
            {"section": 1, "title": "הלכות תלמוד תורה פרק ב'", "author": "רמב\"ם", "kind": "classic",
             "evidence_level": "tradition", "license": "public_domain", "summary_paragraphs": 1},
            {"section": 2, "title": "משלי כ\"ב, ו'", "kind": "classic", "evidence_level": "tradition",
             "license": "public_domain", "summary_paragraphs": 1},
        ],
    }
    return Manifest.model_validate({**data, **changes})


@pytest.fixture
def sample(tmp_path) -> Path:
    document = docx.Document()
    for paragraph in RAMBAM:
        document.add_paragraph(paragraph)
    document.add_paragraph("")  # empty paragraphs are dropped
    path = tmp_path / "sample.docx"
    document.save(str(path))
    return path


# --- extraction and cleaning -----------------------------------------------------------------------


def test_a_word_file_becomes_its_paragraphs_in_order(sample):
    blocks = extract(sample)
    assert [b.text for b in blocks] == RAMBAM and all(b.page is None for b in blocks)


def test_text_files_split_on_blank_lines_and_unknown_types_are_refused(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("פסקה ראשונה\nשורה שנייה שלה\n\nפסקה שנייה", encoding="utf-8")
    assert [b.text for b in extract(path)] == ["פסקה ראשונה\nשורה שנייה שלה", "פסקה שנייה"]
    with pytest.raises(ExtractionError, match="Unsupported"):
        extract(tmp_path / "book.epub")
    (tmp_path / "empty.txt").write_text("  \n\n ", encoding="utf-8")
    with pytest.raises(ExtractionError, match="No text"):
        extract(tmp_path / "empty.txt")


def test_cleaning_removes_what_the_eye_cannot_see_and_joins_broken_lines():
    assert clean_text("מא​יה הלכה\nלמכינה") == "מאיה הלכה למכינה"
    assert clean_text("חינוך מו-\nתאם אישית") == "חינוך מותאם אישית"
    assert clean_text("חֲנֹךְ") == "חֲנֹךְ"  # niqqud is kept in the stored text


def test_page_headers_repeated_on_most_pages_and_page_numbers_are_removed():
    pages = [Block(f"חינוך במכינה\nתוכן של עמוד {n} על הדרכה.\n{n}", page=n) for n in range(1, 5)]
    blocks, removed = clean_blocks(pages)
    assert [b.text for b in blocks] == [f"תוכן של עמוד {n} על הדרכה." for n in range(1, 5)]
    assert "חינוך במכינה" in removed


# --- chunking ---------------------------------------------------------------------------------------


def words(n: int, prefix: str = "מילה") -> str:
    return " ".join(f"{prefix}{i}" for i in range(n))


def test_fixed_windows_overlap_so_nothing_is_lost_at_a_border():
    drafts = fixed([Unit(words(500))], size=200, overlap=40)
    assert [d.words for d in drafts] == [200, 200, 180]
    assert drafts[0].text.split()[-40:] == drafts[1].text.split()[:40]
    with pytest.raises(ValueError):
        fixed([Unit("א")], size=10, overlap=10)


def test_structure_keeps_paragraphs_whole_and_groups_them_up_to_the_limit():
    units = [Unit(words(80, "א"), "הלכה א"), Unit(words(80, "ב"), "הלכה ב"), Unit(words(80, "ג"), "הלכה ג")]
    drafts = structure(units, max_words=200)
    assert [d.words for d in drafts] == [160, 80]
    assert drafts[0].section == "הלכה א, הלכה ב" and drafts[1].section == "הלכה ג"


def test_a_summary_is_a_chunk_of_its_own():
    drafts = structure([Unit("תקציר קצר", SUMMARY), Unit("טקסט המקור")], max_words=200)
    assert [d.text for d in drafts] == ["תקציר קצר", "טקסט המקור"]


def test_a_paragraph_longer_than_the_limit_is_split_at_sentence_ends():
    long = " ".join(f"{words(30, chr(0x05D0 + i))}." for i in range(4))  # four sentences of 30 words
    drafts = structure([Unit(long, "פרק")], max_words=70)
    assert [d.words for d in drafts] == [60, 60]
    assert all(d.text.endswith(".") for d in drafts) and all(d.section == "פרק" for d in drafts)


def test_pages_of_a_chunk_span_its_paragraphs():
    drafts = structure([Unit("א", page=3), Unit("ב", page=4)], max_words=100)
    assert (drafts[0].page_start, drafts[0].page_end) == (3, 4)


@pytest.mark.parametrize(
    ("paragraph", "label"),
    [
        ("הלכה ב: מכניסין את התינוקות", "הלכה ב"),
        ('"הלכה א: מושיבין', "הלכה א"),
        ('פירוש המלבי"ם (מלא): "חנוך לנער"', 'פירוש המלבי"ם (מלא)'),
        ("סתם משפט בלי תווית", None),
        ("משפט ארוך מאוד שיש בו הרבה מאוד מילים לפני הנקודתיים: ואז עוד", None),
    ],
)
def test_labels_at_the_start_of_a_paragraph(paragraph, label):
    assert label_of(paragraph) == label


# --- the manifest and the plan ------------------------------------------------------------------------


def test_numbered_lines_start_sections():
    blocks = [Block("שער"), Block("1. ראשון"), Block("א"), Block("2. שני"), Block("ב")]
    sections = split_sections(blocks, "numbered")
    assert [(s.heading, [b.text for b in s.blocks]) for s in sections] == [("1. ראשון", ["א"]), ("2. שני", ["ב"])]


def test_the_plan_makes_one_source_per_work_with_its_chunks(sample):
    p = plan(sample, manifest())
    rambam, mishlei = p.sources
    assert rambam.entry.title == "הלכות תלמוד תורה פרק ב'"
    assert [d.section for d in rambam.chunks["structure-220"]] == [SUMMARY, "הלכה א, הלכה ב, הלכה ה"]
    assert [d.section for d in mishlei.chunks["structure-220"]] == [SUMMARY, "פסוק"]
    assert "עשרים וחמשה" in preview(p)


def test_a_manifest_that_does_not_fit_the_file_is_refused(sample):
    with pytest.raises(ManifestError, match="not sample.docx"):
        plan(sample, manifest(file="other.docx"))
    too_many = manifest()
    too_many.sources[1].section = 9
    with pytest.raises(ManifestError, match="section 9"):
        plan(sample, too_many)
    with pytest.raises(ManifestError, match="Unknown chunking"):
        plan(sample, manifest(), ["semantic-500"])


def test_the_repo_manifest_for_the_classics_is_valid():
    data = json.loads(Path("library/classics.json").read_text(encoding="utf-8"))
    m = Manifest.model_validate(data)
    assert len(m.sources) == 5 and all(s.evidence_level is EvidenceLevel.TRADITION for s in m.sources)
    eval_set = load_library_set(Path("evals/library-v1.json"))
    assert unknown_sources(eval_set, {s.title for s in m.sources}) == []  # every answer is a real source


# --- import, index, search ----------------------------------------------------------------------------


def test_import_then_import_again_replaces_and_does_not_duplicate(sample):
    uow = FakeUnitOfWork(empty_repos())
    import_file(uow, sample, manifest())
    import_file(uow, sample, manifest())
    with uow() as repos:
        sources = repos.sources.list_all()
        assert [s.title for s in sources] == ["הלכות תלמוד תורה פרק ב'", "משלי כ\"ב, ו'"]
        assert all(s.origin_file == "sample.docx" for s in sources)
        assert len(repos.chunks.list_by_source(sources[0].key, "structure-220")) == 2


def test_the_embedding_input_has_the_work_and_no_niqqud():
    text = for_embedding("חֲנֹךְ לַנַּעַר", title="משלי", section="פסוק")
    # Without niqqud the vocalized spelling stays defective: חנך, not the modern חנוך. One reason
    # meaning-based search (embeddings) suits these texts better than matching words.
    assert text == "משלי · פסוק\nחנך לנער"
    assert for_embedding("שאלה?") == "שאלה?"


def test_the_whole_chain_finds_the_rambam_for_a_question_about_class_size(sample):
    uow = FakeUnitOfWork(empty_repos())
    import_file(uow, sample, manifest())
    embedder = FakeEmbedder()
    index_library(uow, embedder, "structure-220")
    hits = search_library(uow, embedder, "structure-220", "כמה תינוקות למלמד אחד?", k=1)
    assert hits[0].source.title == "הלכות תלמוד תורה פרק ב'"
    assert all("ֲ" not in text for text in embedder.documents)  # no niqqud reached the model
    assert any(text.startswith("משלי") for text in embedder.documents)  # each with its work's name


@pytest.mark.integration
def test_import_index_and_compare_against_postgresql(sample, engine, monkeypatch, tmp_path):
    import pacer.library.__main__ as cli
    from pacer.repositories import repositories

    def uow():
        return repositories(engine)

    monkeypatch.setattr(cli, "_uow", lambda: uow)
    monkeypatch.setattr(cli, "ollama_embedder", lambda name: FakeEmbedder(name))
    import_file(uow, sample, manifest())
    index_library(uow, FakeEmbedder("bge-m3"), "structure-220")
    questions = tmp_path / "q.json"
    cases = [
        {"id": "L01", "question": "כמה תינוקות למלמד אחד?", "kind": "about_source",
         "sources": ["הלכות תלמוד תורה פרק ב'"], "coverage": ["direct"]},
        {"id": "L02", "question": "איך מגדירים את הרשת האלחוטית?", "kind": "off_topic", "sources": [], "coverage": ["none"]},
    ]
    questions.write_text(json.dumps({"name": "t", "version": 1, "about": "test", "cases": cases}, ensure_ascii=False), encoding="utf-8")
    report = cli.compare(questions, ["bge-m3"], ["structure-220"])
    assert "| bge-m3 | structure-220 | 1/1 | 1/1 | 1.00 |" in report and "(t@v1)" in report


def test_a_report_written_to_a_file_reads_right_to_left_in_a_markdown_preview(sample, tmp_path, capsys):
    import pacer.library.__main__ as cli

    out, manifest_path = tmp_path / "preview.md", tmp_path / "manifest.json"
    manifest_path.write_text(manifest().model_dump_json(), encoding="utf-8")
    cli.main(["preview", str(sample), "--manifest", str(manifest_path), "--out", str(out)])
    text = out.read_text(encoding="utf-8")
    assert text.startswith('<div dir="rtl">\n\n#') and text.endswith("\n\n</div>\n")  # blank lines: Markdown renders
